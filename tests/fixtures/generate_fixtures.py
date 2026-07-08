#!/usr/bin/env python3
"""Generate test document fixtures for integration tests."""

import os

from docx import Document
from fpdf import FPDF


def generate_pdf() -> None:
    """Generate sample PDF with headings and paragraphs."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(text="Section One: Access Control")
    pdf.ln()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        w=0,
        text=(
            "Organizations shall implement access control mechanisms that enforce "
            "the principle of least privilege. Access shall be granted based on "
            "roles and responsibilities, with periodic reviews to ensure continued necessity."
        ),
    )
    pdf.ln(5)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(text="Section Two: Audit Logging")
    pdf.ln()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        w=0,
        text=(
            "All access events must be logged and monitored. Audit logs shall include "
            "user identification, timestamp, resource accessed, and action performed. "
            "Logs must be retained for a minimum of 90 days and protected from tampering."
        ),
    )

    output_path = os.path.join(os.path.dirname(__file__), "documents", "sample.pdf")
    pdf.output(output_path)
    print(f"Generated: {output_path}")


def generate_docx() -> None:
    """Generate sample DOCX with headings and paragraphs."""
    doc = Document()
    doc.add_heading("Section One: Access Control", level=1)
    doc.add_paragraph(
        "Organizations shall implement access control mechanisms that enforce "
        "the principle of least privilege. Access shall be granted based on "
        "roles and responsibilities, with periodic reviews to ensure continued necessity."
    )

    doc.add_heading("Section Two: Audit Logging", level=1)
    doc.add_paragraph(
        "All access events must be logged and monitored. Audit logs shall include "
        "user identification, timestamp, resource accessed, and action performed. "
        "Logs must be retained for a minimum of 90 days and protected from tampering."
    )

    output_path = os.path.join(os.path.dirname(__file__), "documents", "sample.docx")
    doc.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    generate_pdf()
    generate_docx()
